"""
Generate professional A4 landscape presentation slides for EyeGaze Inference.
Uses python-pptx for proper PowerPoint output.

Usage:
    python make_slides.py
    open eyegaze_slides.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# -- A4 landscape dimensions
SLIDE_W = Emu(10692000)  # 297mm
SLIDE_H = Emu(7560000)   # 210mm

# -- Colors
BG       = RGBColor(0x0d, 0x11, 0x17)
SURFACE  = RGBColor(0x16, 0x1b, 0x22)
CARD     = RGBColor(0x1c, 0x22, 0x2b)
TEXT     = RGBColor(0xe6, 0xed, 0xf3)
DIM      = RGBColor(0x8b, 0x94, 0x9e)
GREEN    = RGBColor(0x3f, 0xb9, 0x50)
BLUE     = RGBColor(0x58, 0xa6, 0xff)
PURPLE   = RGBColor(0xbc, 0x8c, 0xff)
ORANGE   = RGBColor(0xf0, 0x88, 0x3e)
RED      = RGBColor(0xf8, 0x51, 0x49)
BORDER   = RGBColor(0x30, 0x36, 0x3d)
WHITE    = RGBColor(0xff, 0xff, 0xff)
BLACK    = RGBColor(0x00, 0x00, 0x00)

FONT_TITLE = "Helvetica Neue"
FONT_BODY  = "Helvetica Neue"
FONT_MONO  = "SF Mono"


def _set_slide_bg(slide, color=BG):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_text(slide, left, top, width, height, text, font_size=14,
              color=TEXT, bold=False, font=FONT_BODY, alignment=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font
    p.alignment = alignment
    return txBox


def _add_para(text_frame, text, font_size=12, color=DIM, bold=False,
              font=FONT_BODY, space_before=Pt(4), space_after=Pt(2),
              alignment=PP_ALIGN.LEFT):
    p = text_frame.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font
    p.space_before = space_before
    p.space_after = space_after
    p.alignment = alignment
    return p


def _add_rect(slide, left, top, width, height, fill_color=SURFACE,
              border_color=BORDER, border_width=Pt(1), radius=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = border_color
    shape.line.width = border_width
    if radius is not None:
        shape.adjustments[0] = radius
    return shape


def _add_card(slide, left, top, width, height, title, items,
              accent_color=BLUE, title_size=14, item_size=11):
    """Add a card with title and bullet items."""
    _add_rect(slide, left, top, width, height,
              fill_color=CARD, border_color=accent_color, border_width=Pt(1.5))

    txBox = slide.shapes.add_textbox(
        left + Emu(150000), top + Emu(100000),
        width - Emu(300000), height - Emu(200000))
    tf = txBox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(title_size)
    p.font.color.rgb = accent_color
    p.font.bold = True
    p.font.name = FONT_BODY
    p.space_after = Pt(8)

    for item in items:
        _add_para(tf, item, font_size=item_size, color=DIM, space_before=Pt(3))

    return txBox


def _footer(slide, page_num, total):
    _add_text(slide, Emu(400000), SLIDE_H - Emu(400000), Emu(3000000), Emu(250000),
              "Emanuele Di Luzio", font_size=8, color=BORDER)
    _add_text(slide, SLIDE_W - Emu(3400000), SLIDE_H - Emu(400000),
              Emu(3000000), Emu(250000),
              f"EyeGaze Inference  \u2014  {page_num}/{total}",
              font_size=8, color=BORDER, alignment=PP_ALIGN.RIGHT)


def _title_bar(slide, title):
    _add_text(slide, Emu(400000), Emu(250000), Emu(9000000), Emu(600000),
              title, font_size=28, color=WHITE, bold=True, font=FONT_TITLE)
    # accent line under title
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(400000), Emu(850000), Emu(800000), Emu(35000))
    line.fill.solid()
    line.fill.fore_color.rgb = GREEN
    line.line.fill.background()


def mm(val):
    """Convert mm to Emu."""
    return Emu(int(val * 36000))


# =========================================================================
# SLIDES
# =========================================================================

def slide_title(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    _set_slide_bg(slide)

    # Main title
    _add_text(slide, mm(20), mm(50), mm(257), mm(25),
              "EyeGaze Inference", font_size=48, color=WHITE,
              bold=True, font=FONT_TITLE, alignment=PP_ALIGN.CENTER)

    # Subtitle
    _add_text(slide, mm(20), mm(78), mm(257), mm(15),
              "Real-Time Gaze Estimation with GazeDINO",
              font_size=22, color=DIM, alignment=PP_ALIGN.CENTER)

    # Green line
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, mm(110), mm(98), mm(77), Emu(30000))
    line.fill.solid()
    line.fill.fore_color.rgb = GREEN
    line.line.fill.background()

    # Tech stack
    _add_text(slide, mm(20), mm(105), mm(257), mm(12),
              "DINOv2 ViT-B/14   \u00b7   MediaPipe FaceMesh   \u00b7   Polynomial Calibration",
              font_size=14, color=BLUE, alignment=PP_ALIGN.CENTER)

    # Key metric
    _add_rect(slide, mm(105), mm(125), mm(87), mm(22),
              fill_color=SURFACE, border_color=GREEN, border_width=Pt(2))
    _add_text(slide, mm(105), mm(127), mm(87), mm(18),
              "3.24\u00b0 mean angular error",
              font_size=18, color=GREEN, bold=True, alignment=PP_ALIGN.CENTER)

    # Author
    _add_text(slide, mm(20), mm(165), mm(257), mm(12),
              "Emanuele Di Luzio", font_size=16, color=DIM,
              alignment=PP_ALIGN.CENTER)


def slide_architecture(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _title_bar(slide, "Model Architecture")
    _footer(slide, 2, 7)

    # Left column — pipeline blocks
    blocks = [
        ("Input  (B, 3, 224, 224)", "Face crop, ImageNet-normalised", GREEN),
        ("DINOv2 ViT-B/14", "86M params, self-supervised on 142M images", BLUE),
        ("CLS Token  [768 dim]", "Global image representation", PURPLE),
        ("MLP Head", "LN(768) \u2192 512 (GELU, drop 0.30)\n\u2192 128 (GELU, drop 0.15) \u2192 2", ORANGE),
        ("Output  (yaw, pitch)", "Gaze direction in radians", GREEN),
    ]

    y = mm(30)
    for title, desc, color in blocks:
        _add_rect(slide, mm(10), y, mm(110), mm(24),
                  fill_color=CARD, border_color=color, border_width=Pt(1.5))
        _add_text(slide, mm(14), y + Emu(50000), mm(102), mm(10),
                  title, font_size=12, color=color, bold=True)
        _add_text(slide, mm(14), y + Emu(400000), mm(102), mm(14),
                  desc, font_size=9, color=DIM)
        y += mm(28)

        # arrow between blocks
        if color != GREEN or title.startswith("Input"):
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.DOWN_ARROW, mm(62), y - mm(5), mm(6), mm(5))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = BORDER
            arrow.line.fill.background()

    # Right column — training details
    _add_card(slide, mm(140), mm(30), mm(145), mm(65),
              "Training", [
                  "Dataset: GazeGene (56 subjects, 9 cameras, ~90K samples)",
                  "Loss: Angular error between 3D gaze vectors",
                  "Best validation error: 3.24\u00b0",
                  "50 epochs, AdamW optimizer",
                  "Backbone frozen initially, then fine-tuned",
              ], accent_color=BLUE)

    # Backbone options table
    _add_rect(slide, mm(140), mm(105), mm(145), mm(70),
              fill_color=CARD, border_color=PURPLE, border_width=Pt(1.5))
    _add_text(slide, mm(145), mm(108), mm(100), mm(12),
              "Backbone Options", font_size=14, color=PURPLE, bold=True)

    headers = [("Backbone", mm(145)), ("Params", mm(210)), ("Embed", mm(250))]
    for h, x in headers:
        _add_text(slide, x, mm(122), mm(50), mm(10),
                  h, font_size=10, color=BLUE, bold=True, font=FONT_MONO)

    rows = [
        ("dinov2_vits14", "21M", "384", DIM),
        ("dinov2_vitb14", "86M", "768", GREEN),
        ("dinov2_vitl14", "307M", "1024", DIM),
        ("dinov2_vitg14", "1.1B", "1536", DIM),
    ]
    for i, (name, params, embed, c) in enumerate(rows):
        ry = mm(132 + i * 10)
        _add_text(slide, mm(145), ry, mm(60), mm(10), name, font_size=9, color=c, font=FONT_MONO)
        _add_text(slide, mm(210), ry, mm(35), mm(10), params, font_size=9, color=c, font=FONT_MONO)
        label = embed + "  \u2190 default" if c == GREEN else embed
        _add_text(slide, mm(250), ry, mm(40), mm(10), label, font_size=9, color=c, font=FONT_MONO)


def slide_pipeline(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _title_bar(slide, "Inference Pipeline")
    _footer(slide, 3, 7)

    # Main pipeline flow (left)
    steps = [
        ("Frame (BGR)", "Webcam or video file", GREEN),
        ("MediaPipe FaceMesh", "468 landmarks + iris refinement", BLUE),
        ("Face Crop  224\u00d7224", "Bounding box from landmarks", PURPLE),
        ("GazeDINO", "(yaw, pitch) in radians", ORANGE),
        ("EMA Smoothing", "Stable output, \u03b1 = 0.25", GREEN),
    ]

    y = mm(30)
    for title, desc, color in steps:
        _add_rect(slide, mm(10), y, mm(80), mm(20),
                  fill_color=CARD, border_color=color, border_width=Pt(1.5))
        _add_text(slide, mm(14), y + Emu(50000), mm(72), mm(8),
                  title, font_size=11, color=color, bold=True)
        _add_text(slide, mm(14), y + Emu(350000), mm(72), mm(10),
                  desc, font_size=8, color=DIM)
        y += mm(24)

        if title != "EMA Smoothing":
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.DOWN_ARROW, mm(47), y - mm(5), mm(6), mm(5))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = BORDER
            arrow.line.fill.background()

    # Parallel branches (right side)
    branches = [
        ("Iris & Pupillometry", [
            "Landmarks 468-477 (iris ring)",
            "iris_ratio = iris_diameter / eye_width",
            "EMA smoothed (\u03b1 = 0.2)",
        ], BLUE),
        ("Distance Estimation", [
            "Cheekbone landmarks (234, 454)",
            "d = (14cm \u00d7 focal) / face_px_width",
            "Comfort zone: 40\u201390 cm",
        ], PURPLE),
        ("Head Pose", [
            "6-point solvePnP + Rodrigues",
            "Euler angles: yaw, pitch, roll",
            "3D model: nose, chin, eyes, mouth",
        ], ORANGE),
        ("Blink Detection", [
            "EAR = Eye Aspect Ratio",
            "Vertical / horizontal landmarks",
            "Threshold: 0.21",
        ], RED),
    ]

    y = mm(30)
    for title, items, color in branches:
        _add_card(slide, mm(110), y, mm(80), mm(34),
                  title, items, accent_color=color, title_size=11, item_size=8)
        y += mm(37)

    # Output box
    _add_rect(slide, mm(205), mm(60), mm(80), mm(50),
              fill_color=CARD, border_color=GREEN, border_width=Pt(2))
    _add_text(slide, mm(210), mm(63), mm(70), mm(12),
              "Output", font_size=14, color=GREEN, bold=True)

    outputs = ["Gaze arrows overlay", "Screen (x, y) coords",
               "Distance + status", "Pupil dilation", "Head orientation"]
    txBox = slide.shapes.add_textbox(mm(210), mm(75), mm(70), mm(35))
    tf = txBox.text_frame
    tf.word_wrap = True
    for item in outputs:
        _add_para(tf, "\u2022  " + item, font_size=9, color=DIM, space_before=Pt(2))


def slide_calibration(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _title_bar(slide, "16-Point Gaze Calibration")
    _footer(slide, 4, 7)

    # Left — grid visualization
    grid_x, grid_y = mm(15), mm(35)
    grid_w, grid_h = mm(85), mm(85)
    _add_rect(slide, grid_x, grid_y, grid_w, grid_h,
              fill_color=RGBColor(0x0a, 0x0a, 0x0a), border_color=BORDER, border_width=Pt(2))

    points = [
        (0.08, 0.08), (0.37, 0.08), (0.63, 0.08), (0.92, 0.08),
        (0.08, 0.37), (0.37, 0.37), (0.63, 0.37), (0.92, 0.37),
        (0.08, 0.63), (0.37, 0.63), (0.63, 0.63), (0.92, 0.63),
        (0.08, 0.92), (0.37, 0.92), (0.63, 0.92), (0.92, 0.92),
    ]
    for rx, ry in points:
        cx = grid_x + int(rx * grid_w)
        cy = grid_y + int(ry * grid_h)
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - mm(2), cy - mm(2), mm(4), mm(4))
        dot.fill.solid()
        dot.fill.fore_color.rgb = GREEN
        dot.line.color.rgb = WHITE
        dot.line.width = Pt(1)

    _add_text(slide, mm(15), mm(125), mm(85), mm(12),
              "4\u00d74 grid  \u2022  16 calibration points",
              font_size=10, color=DIM, alignment=PP_ALIGN.CENTER)

    # Right — how it works
    _add_card(slide, mm(115), mm(35), mm(170), mm(55),
              "How It Works", [
                  "1.  Fullscreen overlay shows 16 dots in 4\u00d74 grid",
                  "2.  Follow each dot with your eyes (2.8s per point, ~45s total)",
                  "3.  Sample collection starts after 35% of dwell time",
                  "4.  Outlier filtering: discard samples > 1.5\u03c3 from median",
                  "5.  Degree-3 polynomial least-squares fit",
                  "6.  Minimum 10 valid points required (out of 16)",
                  "7.  Saved to calibration.pkl for reuse across sessions",
              ], accent_color=GREEN, title_size=13, item_size=10)

    # Polynomial features
    _add_rect(slide, mm(115), mm(100), mm(170), mm(38),
              fill_color=CARD, border_color=BLUE, border_width=Pt(1.5))
    _add_text(slide, mm(120), mm(103), mm(160), mm(12),
              "Polynomial Features (degree 3)", font_size=13, color=BLUE, bold=True)
    _add_text(slide, mm(120), mm(115), mm(160), mm(10),
              "[1,  y,  p,  y\u00b2,  yp,  p\u00b2,  y\u00b3,  y\u00b2p,  yp\u00b2,  p\u00b3]",
              font_size=14, color=WHITE, font=FONT_MONO)
    _add_text(slide, mm(120), mm(125), mm(160), mm(10),
              "10 coefficients per axis  \u2022  Two fits: features \u2192 screen_x, features \u2192 screen_y",
              font_size=9, color=DIM)


def slide_screen_gaze(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _title_bar(slide, "Screen Gaze \u2014 Real-Time Overlay")
    _footer(slide, 5, 7)

    # Left — how it works
    _add_card(slide, mm(10), mm(32), mm(130), mm(70),
              "How It Works", [
                  "1.  Run 16-point calibration (~45 seconds)",
                  "2.  Enable \u2018Screen gaze\u2019 toggle in sidebar",
                  "3.  Translucent green dot appears on screen",
                  "4.  Dot follows your gaze position in real-time",
                  "5.  Position smoothed with EMA (\u03b1 = 0.5)",
                  "6.  macOS: truly transparent background (systemTransparent)",
                  "7.  Always-on-top overlay, toggleable on/off",
              ], accent_color=GREEN, title_size=14, item_size=11)

    # Pipeline
    _add_rect(slide, mm(10), mm(110), mm(275), mm(25),
              fill_color=CARD, border_color=BLUE, border_width=Pt(1.5))
    _add_text(slide, mm(15), mm(112), mm(265), mm(10),
              "End-to-End Pipeline", font_size=12, color=BLUE, bold=True)
    _add_text(slide, mm(15), mm(122), mm(265), mm(10),
              "GazeDINO \u2192 (yaw, pitch) \u2192 EMA \u2192 Poly3 Calibrator \u2192 (screen_x, screen_y) \u2192 Position EMA \u2192 Overlay Dot",
              font_size=11, color=DIM, font=FONT_MONO)

    # Right — visual: screen with gaze dot
    scr_x, scr_y = mm(160), mm(32)
    scr_w, scr_h = mm(120), mm(70)
    _add_rect(slide, scr_x, scr_y, scr_w, scr_h,
              fill_color=RGBColor(0x0a, 0x0a, 0x0a), border_color=DIM, border_width=Pt(2))
    _add_text(slide, scr_x, scr_y + Emu(50000), scr_w, mm(8),
              "Screen", font_size=9, color=BORDER, alignment=PP_ALIGN.CENTER)

    # Gaze dot (large green circle)
    dot_x = scr_x + mm(55)
    dot_y = scr_y + mm(35)
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, dot_x, dot_y, mm(12), mm(12))
    dot.fill.solid()
    dot.fill.fore_color.rgb = GREEN
    dot.line.color.rgb = WHITE
    dot.line.width = Pt(2)

    # Trail dots (smaller, dimmer)
    trail = [(-20, 12), (-14, 8), (-8, 5)]
    for dx, dy in trail:
        td = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, dot_x + mm(dx), dot_y + mm(dy), mm(5), mm(5))
        td.fill.solid()
        td.fill.fore_color.rgb = GREEN
        td.line.fill.background()

    _add_text(slide, dot_x - mm(5), dot_y + mm(15), mm(25), mm(8),
              "gaze position", font_size=8, color=GREEN, alignment=PP_ALIGN.CENTER)


def slide_gui(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _title_bar(slide, "GUI & Features")
    _footer(slide, 6, 7)

    # Left — interface
    _add_card(slide, mm(10), mm(32), mm(130), mm(55),
              "Dark-Mode Interface (CustomTkinter)", [
                  "\u2022  Live camera feed with configurable overlays",
                  "\u2022  Sidebar: gaze, metrics, calibration, toggles",
                  "\u2022  Multi-face support (up to 5 simultaneous faces)",
                  "\u2022  Stable face ordering via centroid tracking",
                  "\u2022  Video mode: open files, seek bar, play/pause",
                  "\u2022  1280\u00d7800 default, resizable",
              ], accent_color=BLUE, title_size=12, item_size=10)

    # Overlays
    _add_rect(slide, mm(10), mm(92), mm(60), mm(50),
              fill_color=CARD, border_color=GREEN, border_width=Pt(1.5))
    _add_text(slide, mm(14), mm(95), mm(52), mm(10),
              "Overlay Toggles", font_size=12, color=GREEN, bold=True)
    toggles = ["\u25a0  Face box", "\u25a0  Gaze arrows", "\u25a0  Iris dots",
               "\u25a0  Screen gaze"]
    txBox = slide.shapes.add_textbox(mm(14), mm(107), mm(52), mm(30))
    tf = txBox.text_frame
    tf.word_wrap = True
    for t in toggles:
        _add_para(tf, t, font_size=10, color=DIM, space_before=Pt(3))

    # Shortcuts
    _add_rect(slide, mm(80), mm(92), mm(60), mm(50),
              fill_color=CARD, border_color=ORANGE, border_width=Pt(1.5))
    _add_text(slide, mm(84), mm(95), mm(52), mm(10),
              "Keyboard Shortcuts", font_size=12, color=ORANGE, bold=True)
    shortcuts = ["Q       Quit", "C       Calibrate", "Space   Play/Pause",
                 "\u2190 \u2192     Switch face"]
    txBox = slide.shapes.add_textbox(mm(84), mm(107), mm(52), mm(30))
    tf = txBox.text_frame
    tf.word_wrap = True
    for s in shortcuts:
        _add_para(tf, s, font_size=9, color=DIM, font=FONT_MONO, space_before=Pt(3))

    # Right — live metrics
    _add_card(slide, mm(155), mm(32), mm(130), mm(110),
              "Live Metrics", [], accent_color=PURPLE, title_size=14)

    metrics = [
        ("Yaw / Pitch", "Model output converted to degrees", GREEN),
        ("Screen XY", "After calibration \u2192 pixel coordinates", GREEN),
        ("Distance", "Pinhole model, color-coded 40\u201390 cm zone", BLUE),
        ("Iris ratio", "Pupillometry proxy for dilation", PURPLE),
        ("EAR", "Blink detection, threshold 0.21", ORANGE),
        ("Head pose", "solvePnP \u2192 yaw, pitch, roll in degrees", RED),
    ]
    txBox = slide.shapes.add_textbox(mm(160), mm(50), mm(120), mm(90))
    tf = txBox.text_frame
    tf.word_wrap = True
    for name, desc, color in metrics:
        p = _add_para(tf, name, font_size=11, color=color, bold=True, space_before=Pt(8))
        _add_para(tf, desc, font_size=9, color=DIM, space_before=Pt(1))


def slide_summary(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _title_bar(slide, "Summary")
    _footer(slide, 7, 7)

    # Left — tech stack
    _add_card(slide, mm(10), mm(32), mm(130), mm(60),
              "Tech Stack", [
                  "PyTorch + TorchVision \u2014 DINOv2 backbone, inference",
                  "MediaPipe \u2014 FaceMesh 468 landmarks + iris",
                  "OpenCV \u2014 Video capture, drawing, solvePnP",
                  "CustomTkinter \u2014 Dark-mode GUI",
                  "NumPy \u2014 Polynomial fit, EMA, linear algebra",
              ], accent_color=BLUE, title_size=14, item_size=10)

    # Right — key numbers
    _add_rect(slide, mm(155), mm(32), mm(130), mm(60),
              fill_color=CARD, border_color=GREEN, border_width=Pt(1.5))
    _add_text(slide, mm(160), mm(35), mm(120), mm(12),
              "Key Numbers", font_size=14, color=GREEN, bold=True)

    numbers = [
        ("3.24\u00b0", "Mean angular error"),
        ("86M", "DINOv2 ViT-B/14 parameters"),
        ("16", "Calibration points (4\u00d74 grid)"),
        ("10", "Poly coefficients per axis"),
        ("5", "Max simultaneous faces"),
        ("\u223c30 fps", "Apple Silicon (MPS)"),
    ]
    txBox = slide.shapes.add_textbox(mm(160), mm(50), mm(120), mm(40))
    tf = txBox.text_frame
    tf.word_wrap = True
    for num, desc in numbers:
        p = _add_para(tf, f"{num}    {desc}", font_size=10, color=DIM,
                      space_before=Pt(4))

    # Bottom — full pipeline
    _add_rect(slide, mm(10), mm(102), mm(275), mm(35),
              fill_color=CARD, border_color=BLUE, border_width=Pt(2))
    _add_text(slide, mm(15), mm(105), mm(265), mm(12),
              "End-to-End Pipeline", font_size=16, color=WHITE, bold=True)
    _add_text(slide, mm(15), mm(118), mm(265), mm(15),
              "Webcam  \u2192  MediaPipe FaceMesh  \u2192  Face Crop  \u2192  GazeDINO (yaw, pitch)  \u2192  EMA  \u2192  Poly3 Calibration  \u2192  Screen (x, y)  \u2192  Overlay Dot",
              font_size=12, color=BLUE, font=FONT_MONO)

    # GitHub link
    _add_text(slide, mm(10), mm(150), mm(275), mm(12),
              "github.com/emanuelediluzio/eyegaze-inference",
              font_size=14, color=DIM, alignment=PP_ALIGN.CENTER)


# =========================================================================
# Generate
# =========================================================================

def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slides = [
        ("Title", slide_title),
        ("Architecture", slide_architecture),
        ("Pipeline", slide_pipeline),
        ("Calibration", slide_calibration),
        ("Screen Gaze", slide_screen_gaze),
        ("GUI & Features", slide_gui),
        ("Summary", slide_summary),
    ]

    for i, (name, fn) in enumerate(slides):
        print(f"  Slide {i+1}/{len(slides)}: {name}")
        fn(prs)

    out = "eyegaze_slides.pptx"
    prs.save(out)
    print(f"\nDone! {len(slides)} slides \u2192 {out}")


if __name__ == "__main__":
    main()
